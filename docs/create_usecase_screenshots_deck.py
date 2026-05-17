"""
유즈케이스 스크린샷 덱 - 실제 운영 화면을 슬라이드로
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
import os
from PIL import Image

NAVY = RGBColor(0x0A, 0x1F, 0x44)
CORAL = RGBColor(0xFF, 0x6B, 0x6B)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SCREENS = "/mnt/g/progress/wawa/wawa_smart_erp/docs/screenshots/usecase"
OUT = "/mnt/g/progress/wawa/wawa_smart_erp/docs/WAWA_ERP_UseCase_Screenshots.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "Pretendard"
    return tb


def add_rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def fit_image(slide, path, x, y, max_w, max_h):
    """이미지를 비율 유지하며 max_w/max_h 안에 배치 (중앙 정렬)"""
    with Image.open(path) as im:
        iw, ih = im.size
    ratio = min(max_w / iw, max_h / ih)
    new_w = int(iw * ratio)
    new_h = int(ih * ratio)
    cx = x + (max_w - new_w) // 2
    cy = y + (max_h - new_h) // 2
    slide.shapes.add_picture(path, cx, cy, width=new_w, height=new_h)


def cover_slide():
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, SW, SH, NAVY)
    add_rect(s, 0, Inches(3.0), SW, Inches(0.05), CORAL)
    add_text(s, Inches(0.8), Inches(2.0), Inches(12), Inches(1),
             "WAWA ERP — 유즈케이스 실제 화면", size=44, bold=True, color=WHITE)
    add_text(s, Inches(0.8), Inches(3.2), Inches(12), Inches(0.6),
             "Live Production Screenshots · wawa-smart-erp.pages.dev",
             size=20, color=LIGHT)
    add_text(s, Inches(0.8), Inches(6.6), Inches(12), Inches(0.4),
             "10개 유즈케이스 · 실제 운영 데이터", size=14, color=GRAY)


def screenshot_slide(title, subtitle, img_path, caption_left, caption_right):
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, SW, Inches(0.9), NAVY)
    add_text(s, Inches(0.5), Inches(0.18), Inches(10), Inches(0.4),
             title, size=22, bold=True, color=WHITE)
    add_text(s, Inches(0.5), Inches(0.55), Inches(10), Inches(0.3),
             subtitle, size=12, color=LIGHT)

    # 좌측 캡션
    add_rect(s, Inches(0.4), Inches(1.15), Inches(3.0), Inches(5.8), LIGHT)
    add_text(s, Inches(0.6), Inches(1.3), Inches(2.7), Inches(0.4),
             "사용자 시나리오", size=12, bold=True, color=CORAL)
    add_text(s, Inches(0.6), Inches(1.75), Inches(2.7), Inches(2.5),
             caption_left, size=11, color=NAVY)

    add_text(s, Inches(0.6), Inches(4.3), Inches(2.7), Inches(0.4),
             "핵심 기능", size=12, bold=True, color=CORAL)
    add_text(s, Inches(0.6), Inches(4.75), Inches(2.7), Inches(2.0),
             caption_right, size=11, color=NAVY)

    # 우측 스크린샷 영역
    img_x = Inches(3.7)
    img_y = Inches(1.15)
    img_w = Inches(9.2)
    img_h = Inches(5.8)
    add_rect(s, img_x, img_y, img_w, img_h, WHITE, line=GRAY)
    if os.path.exists(img_path):
        fit_image(s, img_path, img_x + Emu(50000), img_y + Emu(50000),
                  img_w - Emu(100000), img_h - Emu(100000))

    add_text(s, Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.3),
             "wawa-smart-erp.pages.dev · 실제 운영 환경", size=9, color=GRAY)


cover_slide()

screenshot_slide(
    "UC-2 · 로그인 / 멀티 테넌트",
    "학원 선택 → 이름 + PIN 4자리 인증",
    f"{SCREENS}/00-login.png",
    "원장/선생이 아침 출근 후 학원을 선택하고 본인 이름과 PIN으로 로그인.\n\n학원당 데이터는 academy_id로 완전 격리되며, JWT 기반 세션을 발급.",
    "• 학원 슬러그 기반 멀티 테넌트\n• PIN 4자리 (해시 저장)\n• 역할별 접근 (admin/instructor)\n• 7일 JWT 세션",
)

screenshot_slide(
    "UC-1 · 원장 대시보드",
    "오늘의 출결, 보강 현황, 매출 한눈에",
    f"{SCREENS}/uc1-dashboard.png",
    "원장이 출근 직후 가장 먼저 보는 화면. 어제 결석한 학생, 오늘 보강 일정, 미수 학생 목록을 한 화면에서 파악.",
    "• 오늘의 출결 요약\n• 보강 일정 카드\n• 학생 진도 알림\n• 매출 / 미수 현황",
)

screenshot_slide(
    "UC-4 · 타이머 / 출결",
    "체크인 → 자습 시작 → 자동 출석",
    f"{SCREENS}/uc4-timer.png",
    "학생이 학원 도착 후 본인 이름을 누르면 타이머가 시작되고 출석이 자동으로 기록됨.\n\n선생은 실시간 자습 현황을 한눈에 확인.",
    "• 원클릭 체크인\n• 실시간 자습 타이머\n• 자동 출결 기록\n• 일일/월간 학습 시간 집계",
)

screenshot_slide(
    "UC-5 · 보드 / 할일",
    "학생별 오늘의 할일 + 진행 체크",
    f"{SCREENS}/uc5-board.png",
    "선생이 학생별로 오늘 풀어야 할 문제집·과제를 보드에 등록하고, 학생은 완료 시 체크.\n\n포모도로 타이머와 연동.",
    "• 학생별 할일 카드\n• 드래그 정렬\n• 진행률 자동 집계\n• 모바일 최적화",
)

screenshot_slide(
    "UC-6 · 학생 관리",
    "선생 중심 학생 목록 + 프로필",
    f"{SCREENS}/uc6-students.png",
    "본인이 담당하는 학생만 보임 (admin은 전체).\n\n학생 카드 클릭 → 출결, 성적, 리포트 히스토리 한 페이지에서 확인.",
    "• 담당 학생 자동 필터\n• 학년/학교 검색\n• 출결·성적·리포트 통합 뷰\n• 학부모 연락처 / 메모",
)

screenshot_slide(
    "UC-7 · AI 리포트",
    "Gemini 기반 월간 리포트 자동 생성",
    f"{SCREENS}/uc7-reports.png",
    "월말에 학생별 출결·성적·자습시간 데이터를 Gemini API로 분석.\n\n학부모용 자연어 리포트가 자동 생성되고 알림톡으로 발송.",
    "• Gemini 1.5 기반 분석\n• 학생별 맞춤 코멘트\n• 학부모 공유 링크\n• 알림톡 자동 발송",
)

screenshot_slide(
    "UC-8 · 결석 / 보강 관리",
    "결석 등록 → 보강 일정 → 완료 처리",
    f"{SCREENS}/uc8-absence.png",
    "학부모가 결석 통보 시 선생이 결석 등록.\n\n시스템이 자동으로 보강 후보 시간대를 제안하고 일정 확정 시 학부모에게 알림.",
    "• 결석 사유 기록\n• 보강 일정 자동 매칭\n• 완료 처리 워크플로우\n• 학부모 알림 자동화",
)

screenshot_slide(
    "★ 핵심 · 보강 관리 전체 뷰",
    "미보강 / 보강예정 / 보강완료 한 화면에서",
    f"{SCREENS}/uc-makeup-overview.png",
    "보강은 학원 운영의 가장 큰 페인포인트.\n\n결석 → 보강 매칭 → 완료까지 누락 없이 추적되어야 매출과 학부모 신뢰가 유지됨.",
    "• 상태별 필터 (미보강/예정/완료)\n• 학생별 누적 결석 카운트\n• 보강 일정 즉시 지정\n• 완료 처리 1클릭",
)

screenshot_slide(
    "★ 핵심 · 미보강 추적",
    "결석 후 아직 보강 일정이 잡히지 않은 케이스",
    f"{SCREENS}/uc-makeup-pending.png",
    "결석은 기록됐지만 보강 일정이 미정인 학생을 따로 모아 보여줌.\n\n원장이 매주 이 화면을 점검해 누락된 보강이 없도록 관리.",
    "• 미보강 자동 집계\n• 학부모별 일정 협의\n• 보강일 지정 즉시 상태 전환\n• 누락 방지 = 매출 보호",
)

screenshot_slide(
    "★ 핵심 · 보강예정 / 완료 처리",
    "확정된 보강을 추적하고 완료 시 1클릭",
    f"{SCREENS}/uc-makeup-scheduled.png",
    "보강일이 확정된 학생 목록.\n\n선생은 보강이 끝나면 완료 버튼만 누르면 학부모 알림과 출결이 자동 처리됨.",
    "• 보강 예정 일자별 정렬\n• 학부모 자동 알림\n• 완료 시 출결 자동 기록\n• AI 리포트에 보강 내역 반영",
)

screenshot_slide(
    "UC-10 · 가챠 게이미피케이션",
    "문제 풀이 → 카드 획득 → 컬렉션",
    f"{SCREENS}/uc10-gacha.png",
    "학생이 카드(문제)를 풀면 정답 시 신규 카드를 획득.\n\n카드 컬렉션 시스템으로 학습 동기를 자연스럽게 부여.",
    "• 텍스트 / 이미지 카드\n• 정답 시 카드 획득\n• 컬렉션 도감\n• 선생이 직접 카드 출제",
)

screenshot_slide(
    "UC · 시험 관리",
    "정기고사 / 월말평가 자동 생성",
    f"{SCREENS}/uc-exam-mgmt.png",
    "정기고사·월말평가 설정 시 과목별 시험 인스턴스가 자동 생성.\n\n성적 입력 페이지로 바로 연결되어 리포트와 자동 연동.",
    "• 정기고사 / 월말 / 단원\n• 과목별 자동 분리\n• 성적 입력 즉시 반영\n• AI 리포트 연동",
)

screenshot_slide(
    "UC-2 · 학원 관리 / 초대",
    "원장 → 선생 초대 → 역할 부여",
    f"{SCREENS}/uc2-academy.png",
    "원장이 새 선생을 초대하면 슬러그 + PIN으로 가입.\n\n역할별 권한이 자동 부여되고, 데이터는 학원 단위로 완전 격리.",
    "• 학원 슬러그 / 정보 관리\n• 선생 초대 / 역할 부여\n• 학원별 설정 (시간대, 알림)\n• 멀티 테넌트 격리",
)


prs.save(OUT)
print(f"✅ {OUT}")
print(f"   슬라이드: {len(prs.slides)}장")
